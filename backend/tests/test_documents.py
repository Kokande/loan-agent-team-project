from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from src.rag.documents import Page, chunk_pages, load_pages


def test_chunks_preserve_text_and_overlap():
    chunks = chunk_pages([Page("abcdefghijklmnopqrst", 7)], chunk_size=10, overlap=3)
    assert [chunk.text for chunk in chunks] == ["abcdefghij", "hijklmnopq", "opqrst"]
    assert [chunk.index for chunk in chunks] == [0, 1, 2]
    assert all(chunk.page == 7 for chunk in chunks)


def test_chunks_prefer_paragraph_boundary():
    chunks = chunk_pages([Page("Первый пункт.\n\nВторой пункт длиннее.")], 25, 0)
    assert chunks[0].text == "Первый пункт."
    assert "Второй пункт длиннее." in [chunk.text for chunk in chunks]


def test_pdf_pages_are_not_merged():
    chunks = chunk_pages([Page("Страница один", 1), Page("Страница три", 3)], 100, 10)
    assert [(chunk.text, chunk.page) for chunk in chunks] == [
        ("Страница один", 1), ("Страница три", 3),
    ]


@pytest.mark.parametrize("size,overlap", [(0, 0), (10, -1), (10, 10), (10, 11)])
def test_invalid_chunk_settings_rejected(size, overlap):
    with pytest.raises(ValueError):
        chunk_pages([Page("Текст")], size, overlap)


def test_load_utf8_preserves_paragraphs_and_removes_nulls(tmp_path):
    path = tmp_path / "regulation.md"
    path.write_text("\ufeffПункт 1.\x00\r\n\r\nПункт 2.", encoding="utf-8")
    assert load_pages(path) == [Page("Пункт 1.\n\nПункт 2.")]


def test_load_docx_includes_tables_in_document_order(tmp_path):
    document = Document()
    document.add_paragraph("До таблицы")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "Значение"
    document.add_paragraph("После таблицы")
    path = tmp_path / "regulation.docx"
    document.save(path)
    text = load_pages(path)[0].text
    assert text.index("До таблицы") < text.index("Показатель") < text.index("После таблицы")
    assert "Значение" in text


def test_empty_pdf_requires_ocr_or_text(tmp_path):
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    with pytest.raises(ValueError, match="OCR"):
        load_pages(path)


def test_empty_text_rejected(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text(" \n\t")
    with pytest.raises(ValueError, match="текст"):
        load_pages(path)


def test_unsupported_format_rejected():
    with pytest.raises(ValueError, match="формат"):
        load_pages(Path("regulation.exe"))
